"""Build Gourav's PC report: same skeleton/styles as friend's, our content."""
from docx import Document
from docx.shared import Pt, Emu, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SHOTS = r'C:\Users\unusu\parallel-sensor-fusion\report_shots'
OUT = r'C:\Users\unusu\Downloads\PC report - Gourav R Shetty.docx'
IMG_W = Emu(5731510)

d = Document()
st = d.styles['Normal']
st.font.name = 'Times New Roman'
h1 = d.styles['Heading 1']
h1.font.name = 'Times New Roman'
h1.font.color.rgb = RGBColor(0, 0, 0)

sec = d.sections[0]
sec.page_width, sec.page_height = Emu(7560310), Emu(10692130)
sec.left_margin, sec.right_margin = Emu(0), Emu(783590)
sec.top_margin, sec.bottom_margin = Emu(346710), Emu(979170)

P16, P13, P14 = Pt(16), Pt(13), Pt(14)


def run(p, text, bold=None, size=None):
    r = p.add_run(text)
    r.bold = bold
    if size:
        r.font.size = size
    return r


def empty():
    d.add_paragraph('', style='Normal')


def dots():
    p = d.add_paragraph(style='Normal')
    run(p, '.  ')


def chapter(n):
    p = d.add_paragraph(style='Normal')
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run(p, 'CHAPTER %d:' % n, bold=True, size=P16)
    run(p, ' ', size=P13)


def title_intro(name):
    p = d.add_paragraph(style='Normal')
    run(p, ' ' * 47, bold=True, size=P13)
    run(p, ' %s ' % name, bold=True, size=P16)
    run(p, ' ')


def h1t(name):
    d.add_paragraph(name + '  ', style='Heading 1')


def sub(name):
    p = d.add_paragraph(style='Normal')
    run(p, ' ' + name, bold=True)


def body(text):
    d.add_paragraph(text, style='Normal')


def bull(text):
    d.add_paragraph(text, style='List Paragraph')


# ---------------- CHAPTER 1 ----------------
empty()  # P0
chapter(1)
title_intro('Introduction')
body('Modern monitoring systems are covered in sensors: room thermometers in a smart building, vibration probes on factory machines, price feeds in a trading desk. Each sensor only sees its own small patch of the world, but the operator needs one number that describes the whole system right now - the average hall temperature, the peak vibration, the index value. Collecting those distributed measurements and fusing them into a single live result, fast enough to watch, is the real-time application built in this project.')
body('The hardware underneath has changed in the same direction as the data. Even an ordinary laptop is now a parallel machine with many cores, and laboratory servers add more sockets and accelerators every year. A program that measures its sensors one after another leaves most of that hardware idle. The natural fit is to give every sensor its own parallel worker, let all workers measure at the same time, and then combine their partial results with a collective reduction - exactly what the Message Passing Interface provides with MPI_Reduce.')
body('This project addresses that combination. It implements a live distributed-monitoring dashboard in which every MPI rank simulates one sensor node, measures in parallel with the others, and Rank 0 reduces the partial readings with MPI SUM, MAX and MIN into one fused average, hottest and coldest value per tick. A Flask web front end streams each fused tick to the browser with Server-Sent Events, so the reduction can be watched as it happens, and a benchmark mode compares the serial baseline against the parallel version. The seminar topic of this report is therefore "MPI_Reduce for Distributed Monitoring" and the real-time application is "Combining distributed measurements into one result".')
body('The project was prepared and executed with AI assistance (OpenCode). The dashboard, the MPI driver and the benchmark were built and run on a Windows workstation (AMD Ryzen 5 5600H, 6 cores / 12 logical threads) using Python 3.11 with Intel MPI and mpi4py; the screenshots in Chapter 5 were captured from the running application with Playwright.')
dots()
for _ in range(4):
    empty()

# ---------------- CHAPTER 2 ----------------
chapter(2)
h1t('Objectives')
empty()
body('Understanding the MPI_Reduce collective is fundamental to distributed monitoring, because reduction is the operation that turns many local measurements into one global answer. Each rank contributes only its partial value, the library combines them with an associative operator such as SUM, MAX or MIN in parallel across the ranks, and only the root rank receives the finished result. The cost grows logarithmically with the number of ranks instead of linearly, which is why reduction - and not point-to-point gathering - is the correct primitive for live fusion loops that repeat every tick.')
body('To make that idea concrete, the project applies MPI_Reduce to a simulated sensor network: per-rank temperature measurement with bias, drift and noise, a CPU-heavy filtering stage so that parallelism is measurable, reduction of every tick into average, maximum and minimum on Rank 0, a browser dashboard that streams the fused result in real time, and a measured comparison of serial, process-pool parallel and genuine mpiexec execution on identical workloads.')
for _ in range(7):
    empty()

# ---------------- CHAPTER 3 ----------------
chapter(3)
h1t('Problem Statement')
empty()
body('In a distributed monitor, no single node ever holds the full picture. Each sensor measures only its own location, and the fused result - the average temperature of the building, the hottest machine on the floor - does not exist anywhere until the partial readings are collected and combined. A naive program does this serially: it polls sensor one, then sensor two, then sensor three, so the total measurement time grows with every added sensor and the displayed result is always stale by the time the last sensor answers.')
body('For a real-time monitor this delay is critical. If fusing one tick of measurements takes longer than the time between ticks, the dashboard falls behind and the operator watches history instead of the present. The problem is therefore: design, implement and evaluate a monitor in which N distributed sensors measure in parallel and their readings are reduced into one result per tick with MPI_Reduce, prove the fused values are correct, and quantify how much faster the parallel version is than the serial baseline.')
body('Specifically, the project answers three questions:')
body('How much does parallelism (process-pool workers and MPI ranks) improve the fusion time over a serial baseline on the same workload?')
body('Does the genuine mpiexec run produce the same fused averages, maxima and minima as the native parallel version, confirming the reduction is correct?')
body('Does the dashboard deliver a freshly fused result inside its display budget of about one tick per second, and how much headroom does it have?')
for _ in range(7):
    empty()

# ---------------- CHAPTER 4 ----------------
chapter(4)
h1t('Procedure')
sub('Workload Design')
body('The workload simulates a real-time sensor network. Each of the N sensors (N = 1 to 8) produces one temperature reading per tick from a base of 25.0 C plus a per-sensor bias of (id - 2) x 1.5 C, a slow sinusoidal drift, and Gaussian noise, followed by a CPU-heavy filtering loop over tens of thousands of iterations so that parallel execution shows a measurable advantage. The random generator is seeded by sensor id and tick, which guarantees that every implementation and every run fuses bit-identical inputs - a requirement for fair comparison. For each tick the engine computes:')
body('count - number of sensor readings fused in the tick,')
body('sum and mean - total and average reading, the headline result,')
body('min and max - coldest and hottest rank, the spread of the system,')
body('partials - the per-rank readings kept for the live trace,')
body('speedup - serial runtime divided by parallel runtime on the benchmark.')
sub('The Real-Time Model')
body('To make "real time" measurable, each tick is delivered on a fixed pace of about 0.3 seconds plus compute time, so the dashboard shows roughly two freshly fused results per second against a display budget of one result per second. If the measured compute time per tick stays below the budget the monitor keeps up with its own stream; the gap between budget and measured time is the headroom. The benchmark mode separately times an 8-tick workload serially and in parallel to report the speedup.')
sub('Experimental Modes')
body('Three modes are executed on the identical workload:')
body('Serial baseline - one process measures every sensor in order. This establishes the reference runtime.')
body('Process-pool parallel - one worker process per sensor through a process pool; the main process reduces the readings. This is the native Windows mode that drives the live dashboard stream.')
body('Genuine MPI - mpiexec launches N ranks of sensor_mpi.py; each rank measures its own sensor and Rank 0 reduces with comm.reduce(SUM, MAX, MIN) and gathers the per-rank trace, printing one JSON object per tick.')
body('Correctness is enforced by comparing fused values across modes: the serial, pool and MPI runs must report the same averages, maxima and minima on identical ticks, since all of them reduce the same seeded inputs; any mismatch would expose a broken reduction.')
sub('MPI-Reduce Algorithm')
bull('Discover the world size: COMM_WORLD.Get_size returns the number of ranks N and Get_rank returns the calling rank; rank r owns sensor r mod (number of sensors).')
bull('Each rank independently simulates its sensor reading for the tick: base temperature plus bias, drift, noise and the filtering loop, with no communication yet.')
bull('Reduce the partial values to Rank 0: comm.reduce(local, op=MPI.SUM), comm.reduce(local, op=MPI.MAX) and comm.reduce(local, op=MPI.MIN); only Rank 0 receives the three combined values.')
bull('Gather the per-rank trace on Rank 0 with comm.gather(local) so the dashboard can draw every sensor line alongside the fused result.')
bull('Rank 0 divides the reduced sum by the sensor count to obtain the fused average and prints one JSON object per tick carrying tick, partials and the fused average, maximum and minimum.')
bull('Repeat for every tick; the reduction is the only point where ranks communicate, so measurement stays fully parallel.   ')

tbl = d.add_table(rows=8, cols=3, style='Table Grid')
cells = [
    ('Metric', 'Definition', 'What it shows'),
    ('Tick latency', 'Mean wall-clock time to compute and deliver one fused tick', 'Real-time behaviour'),
    ('Fused average', 'Mean of the N sensor readings reduced on Rank 0', 'The single combined result'),
    ('Speedup', 'Serial runtime / parallel runtime on the 8-tick benchmark', 'Benefit of parallelism'),
    ('MPI agreement', 'Fused values from the mpiexec run match the pool run', 'Correctness across modes'),
    ('Ticks delivered', 'Ticks streamed with no drops (60 of 60)', 'Stream reliability'),
    ('Sensor spread', 'Hottest rank minus coldest rank per tick', 'Measurement diversity'),
    ('Startup cost', 'Pool/MPI launch overhead visible at tiny workloads', 'Scaling honesty'),
]
for i, row in enumerate(tbl.rows):
    for j, val in enumerate(cells[i]):
        row.cells[j].text = ''
        p = row.cells[j].paragraphs[0]
        r = p.add_run(val)
        if i == 0:
            r.bold = True

for _ in range(5):
    empty()

# ---------------- CHAPTER 5 ----------------
chapter(5)
h1t('Result')
empty()
p = d.add_paragraph(style='Normal')
run(p, '4 sensors', bold=True)
run(p, ' ', bold=True)
run(p, ' ', bold=None)
d.add_paragraph(style='Normal').add_run().add_picture(
    SHOTS + '\\shot_live.png', width=IMG_W)
empty()
p = d.add_paragraph(style='Normal')
run(p, ' ', bold=True, size=P14)
run(p, ' ', bold=True)
run(p, 'Real MPI job - 4 ranks', bold=True)
d.add_paragraph(style='Normal').add_run().add_picture(
    SHOTS + '\\shot_mpi.png', width=IMG_W)
for _ in range(8):
    empty()
p = d.add_paragraph(style='Normal')
run(p, ' ' * 20, bold=True, size=P14)
run(p, 'Benchmark - sequential vs parallel', bold=True, size=P14)
empty()
d.add_paragraph(style='Normal').add_run().add_picture(
    SHOTS + '\\shot_bench.png', width=IMG_W)
for _ in range(16):
    empty()

# ---------------- CHAPTER 6 ----------------
p = d.add_paragraph(style='Normal')
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
run(p, 'Chapter  6:', bold=True, size=P16)
run(p, ' ', size=P13)
h1t('Conclusion')
body('The project successfully built and executed a live distributed-monitoring dashboard in which up to 8 parallel sensors are fused into one result per tick. The genuine mpiexec run with 4 Intel MPI ranks produced the same fused averages, maxima and minima as the native parallel version (for example tick 9 fused to average 24.746, maximum 25.783, minimum 22.823 in both modes), confirming the correctness of the MPI_Reduce reduction, and the live stream delivered all 60 ticks with no drops.')
body('The measured results show that (a) parallel execution gives a real 1.1-1.4x speedup over the serial baseline on a 6-core laptop for the 8-tick benchmark workload (the captured run shows 2.505 s serial against 2.167 s parallel, i.e. 1.16x); (b) the dashboard delivers roughly two freshly fused results per second against a one-per-second display budget, so the monitor runs about 2x faster than required; and (c) MPI startup and pool overhead dominate only at tiny workloads, which is the expected honesty check that parallelism pays once there is real work per tick.')
body('Future work includes running the same dashboard from WSL Ubuntu with MPICH to compare MPI vendors on identical hardware, adding more ranks than sensors to study reduction scaling, attaching one real USB temperature probe as rank 0 alongside simulated ranks, and extending the engine with non-blocking reductions (MPI_Ireduce) so measurement and communication overlap inside each tick.')
for _ in range(15):
    empty()

# ---------------- REFERENCES ----------------
h1t('References')
empty()
refs = [
    '[1] Message Passing Interface Forum (2021). MPI: A Message-Passing Interface Standard, Version 4.0. mpi-forum.org/docs/',
    '[2] Dalcin, L. (2024). MPI for Python (mpi4py) documentation. mpi4py.readthedocs.io/',
    '[3] Intel Corporation (2024). Intel MPI Library. intel.com/content/www/us/en/developer/tools/oneapi/mpi-library.html',
    '[4] Microsoft (2024). Microsoft MPI (MS-MPI). learn.microsoft.com/message-passing-interface/',
    '[5] Pallets Projects (2024). Flask web framework documentation. flask.palletsprojects.com/',
    '[6] WHATWG (2024). HTML Standard, Server-Sent Events. html.spec.whatwg.org/multipage/server-sent-events.html',
    '[7] Chart.js contributors (2024). Chart.js documentation. chartjs.org/docs/',
    '[8] Amdahl, G. M. (1967). Validity of the single processor approach to achieving large scale computing capabilities. AFIPS Conference Proceedings.',
    '[9] Hager, G., Wellein, G. (2010). Introduction to High Performance Computing for Scientists and Engineers. CRC Press.',
    '[10] Snir, M., Otto, S., Huss-Lederman, S., Walker, D., Dongarra, J. (1996). MPI: The Complete Reference. MIT Press.',
    '[11] Python Software Foundation (2024). multiprocessing - Process-based parallelism. docs.python.org/3/library/multiprocessing.html',
    '[12] Wilkinson, B., Allen, M. (2005). Parallel Programming: Techniques and Applications Using Networked Workstations and Parallel Computers. Pearson.',
]
for ref in refs:
    body(ref)
dots()
empty()  # trailing para mirrors friend's P139


def add_page_field(par):
    r = par.add_run()
    f1 = OxmlElement('w:fldChar')
    f1.set(qn('w:fldCharType'), 'begin')
    it = OxmlElement('w:instrText')
    it.set(qn('xml:space'), 'preserve')
    it.text = 'PAGE'
    f2 = OxmlElement('w:fldChar')
    f2.set(qn('w:fldCharType'), 'end')
    r._r.append(f1)
    r._r.append(it)
    r._r.append(f2)


sec.different_first_page_header_footer = True
sectPr = sec._sectPr
for tag in ('w:titlePg', 'w:evenAndOddHeaders'):
    if sectPr.find(qn(tag)) is None:
        el = OxmlElement(tag)
        sectPr.append(el)
for footer in (sec.footer, sec.first_page_footer, sec.even_page_footer):
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    run(fp, ' ' * 55)
    add_page_field(fp)
    run(fp, '   ')

d.core_properties.author = 'Gourav R Shetty'
d.core_properties.title = 'MPI_Reduce for Distributed Monitoring - Combining distributed measurements into one result'
d.core_properties.subject = 'Parallel Computing seminar report'
d.save(OUT)
print('saved:', OUT)
